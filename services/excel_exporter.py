import logging
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from models.card import Card

logger = logging.getLogger(__name__)


class CollectionExporter:

    @staticmethod
    def export_to_excel(cards: list[Card], filename: str = "") -> tuple[bool, str]:
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Kolekcija Kartica"

            headers = [
                "Naziv Kartice", "Set", "Tip", "CMC", "Boje",
                "Raritet", "Kolekcijski Broj", "Foil", "Količina", "Cena (€)", "Cena (RSD)",
            ]
            ws.append(headers)

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")
            
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center

            start_row = 2
            for card in cards:
                foil_text = "Da" if card.is_foil else "Ne"
                quantity = getattr(card, "quantity", getattr(card, "count", 1))

                total_card_eur = card.price_numeric * quantity
                total_card_rsd = card.price_rsd * quantity

                ws.append([
                    card.name,
                    card.set_name,
                    card.type_line,
                    card.cmc,
                    card.colors,
                    card.rarity,
                    card.collector_number,
                    foil_text,
                    quantity,
                    total_card_eur,
                    total_card_rsd,  
                ])

            end_row = len(cards) + 1
            total_row_idx = end_row + 2

            total_label = ws.cell(row=total_row_idx, column=8, value="UKUPNO:")
            total_label.font = Font(name="Calibri", size=11, bold=True)
            total_label.alignment = Alignment(horizontal="right")

            total_cards_cell = ws.cell(row=total_row_idx, column=9, value=f"=SUM(I{start_row}:I{end_row})")
            total_cards_cell.font = Font(name="Calibri", size=11, bold=True)
            total_cards_cell.alignment = Alignment(horizontal="center")

            total_eur_cell = ws.cell(row=total_row_idx, column=10, value=f"=SUM(J{start_row}:J{end_row})")
            total_eur_cell.font = Font(name="Calibri", size=11, bold=True, color="1B5E20")
            total_eur_cell.number_format = '€#,##0.00'

            total_rsd_cell = ws.cell(row=total_row_idx, column=11, value=f"=SUM(K{start_row}:K{end_row})")
            total_rsd_cell.font = Font(name="Calibri", size=11, bold=True, color="1B5E20")
            total_rsd_cell.number_format = '#,##0 "RSD"'

            for row in range(start_row, end_row + 1):
                ws.cell(row=row, column=2).alignment = align_center
                ws.cell(row=row, column=10).number_format = '€#,##0.00'
                ws.cell(row=row, column=11).number_format = '#,##0 "RSD"'

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(filename)
            return True, os.path.abspath(filename)

        except Exception as exc:
            logger.error(f"Greška pri eksportu u Excel: {exc}")
            return False, str(exc)