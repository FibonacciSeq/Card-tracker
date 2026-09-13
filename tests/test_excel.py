import openpyxl
import pytest

from models.card import Card
from services.card_collection import CardCollection
from services.excel_exporter import CollectionExporter
from services.excel_importer import ExcelImporter


def _export(cards, path):
    ok, _ = CollectionExporter.export_to_excel(cards, str(path))
    assert ok is True
    return path


def test_export_writes_a_readable_workbook(tmp_path, bolt_data, forest_data):
    path = _export([Card(bolt_data), Card(forest_data)], tmp_path / "col.xlsx")

    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))

    assert rows[0][0] == "Naziv Kartice"
    assert [r[0] for r in rows[1:3]] == ["Lightning Bolt", "Forest"]


def test_export_marks_foil_status(tmp_path, bolt_data):
    path = _export([Card(bolt_data, is_foil=True), Card(bolt_data, is_foil=False)], tmp_path / "col.xlsx")

    ws = openpyxl.load_workbook(path).active
    foil_col = [r[7] for r in list(ws.iter_rows(values_only=True))[1:3]]
    assert foil_col == ["Da", "Ne"]


def test_export_multiplies_price_by_quantity(tmp_path, bolt_data):
    path = _export([Card(bolt_data).clone(quantity=3)], tmp_path / "col.xlsx")

    ws = openpyxl.load_workbook(path).active
    row = list(ws.iter_rows(values_only=True))[1]
    assert row[8] == 3
    assert row[9] == pytest.approx(13.50)


def test_export_adds_a_totals_row(tmp_path, bolt_data):
    path = _export([Card(bolt_data)], tmp_path / "col.xlsx")

    ws = openpyxl.load_workbook(path).active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "UKUPNO:" in values
    assert any(isinstance(v, str) and v.startswith("=SUM(") for v in values)


def test_export_of_an_empty_collection_still_succeeds(tmp_path):
    ok, _ = CollectionExporter.export_to_excel([], str(tmp_path / "empty.xlsx"))
    assert ok is True


def test_export_import_roundtrip(tmp_path, fake_db, bolt_data, forest_data):
    path = _export(
        [Card(bolt_data).clone(quantity=2), Card(forest_data).clone(quantity=4)],
        tmp_path / "col.xlsx",
    )

    collection = CardCollection()
    added, unmatched = ExcelImporter.import_file(str(path), fake_db, collection)

    assert added == 6
    assert unmatched == []
    assert collection.get_total_count() == 6


def test_import_reports_unknown_cards(tmp_path, fake_db):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Naziv", "Kolicina"])
    ws.append(["Black Lotus", 1])
    ws.append(["Forest", 2])
    path = tmp_path / "mixed.xlsx"
    wb.save(path)

    collection = CardCollection()
    added, unmatched = ExcelImporter.import_file(str(path), fake_db, collection)

    assert added == 2
    assert unmatched == ["1x Black Lotus"]


def test_import_skips_rows_without_a_name(tmp_path, fake_db):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Naziv", "Kolicina"])
    ws.append([None, 3])
    ws.append(["Forest", 1])
    path = tmp_path / "gaps.xlsx"
    wb.save(path)

    collection = CardCollection()
    added, _ = ExcelImporter.import_file(str(path), fake_db, collection)
    assert added == 1


def test_import_of_an_empty_sheet_is_harmless(tmp_path, fake_db):
    wb = openpyxl.Workbook()
    path = tmp_path / "blank.xlsx"
    wb.save(path)

    assert ExcelImporter.import_file(str(path), fake_db, CardCollection()) == (0, [])


def _sheet(tmp_path, rows, name="sheet.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


def test_import_handles_a_leading_blank_column(tmp_path, fake_db):
    """Regresija: prazna prva kolona je pomerala indekse i uvozila 0 karata."""
    path = _sheet(tmp_path, [[None, "Naziv", "Kolicina"], [None, "Forest", 4]])

    collection = CardCollection()
    added, unmatched = ExcelImporter.import_file(str(path), fake_db, collection)

    assert (added, unmatched) == (4, [])
    assert collection.get_total_count() == 4


def test_import_handles_blank_columns_between_headers(tmp_path, fake_db):
    path = _sheet(tmp_path, [["Naziv", None, "Kolicina"], ["Forest", None, 2]])

    collection = CardCollection()
    added, _ = ExcelImporter.import_file(str(path), fake_db, collection)
    assert added == 2


def test_first_matching_header_wins(tmp_path, fake_db):
    """'Card Name' i 'Card #' oba sadrze 'card' - ime je prva kolona."""
    path = _sheet(tmp_path, [["Card Name", "Card #"], ["Forest", "280"]])

    collection = CardCollection()
    added, _ = ExcelImporter.import_file(str(path), fake_db, collection)
    assert added == 1
    assert collection.items[0].name == "Forest"
