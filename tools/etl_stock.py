"""Korak 2 posla: Excel/CSV iz drop foldera -> staging -> inv.Stock.

Pokretanje:
    python -m tools.etl_stock --database CardTracker --drop-folder C:\\CardTracker\\drop

Obradjeni fajlovi se premestaju u podfolder 'archive', a oni koji nisu mogli
da se procitaju u 'failed', da posao sledeci put ne bi ponovo naleteo na njih.
"""

import argparse
import csv
import logging
import os
import shutil
import sys
from datetime import datetime

import openpyxl

from services.mssql import ConnectionSettings, connection, sql
from services.spreadsheet import cell, detect_columns

logger = logging.getLogger("etl.stock")

STAGING_COLUMNS = (
    "SourceRowNo", "CardName", "SetCode", "CollectorNumber", "Quantity",
    "IsFoil", "ConditionCode", "LanguageCode", "Location", "UnitCostRsd",
)

FIELD_ORDER = (
    "name", "set_code", "collector_number", "quantity",
    "is_foil", "condition", "language", "location", "unit_cost",
)


def read_rows(path: str) -> list[tuple]:
    """Ucitava .xlsx ili .csv u listu redova."""
    extension = os.path.splitext(path)[1].lower()

    if extension in (".xlsx", ".xlsm"):
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            return list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()

    if extension in (".csv", ".tsv"):
        delimiter = "\t" if extension == ".tsv" else ","
        with open(path, encoding="utf-8-sig", newline="") as handle:
            return [tuple(row) for row in csv.reader(handle, delimiter=delimiter)]

    raise ValueError(f"Nepodržan format fajla: {extension}")


def rows_to_staging(rows: list[tuple]) -> list[tuple]:
    """Redovi tabele -> torke spremne za staging.StockImport."""
    start_row, columns = detect_columns(rows)

    staged = []
    for offset, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row:
            continue

        name = cell(row, columns, "name")
        if not name:
            continue

        staged.append((
            offset,
            name,
            cell(row, columns, "set_code"),
            cell(row, columns, "collector_number"),
            cell(row, columns, "quantity"),
            cell(row, columns, "is_foil"),
            cell(row, columns, "condition"),
            cell(row, columns, "language"),
            cell(row, columns, "location"),
            cell(row, columns, "unit_cost"),
        ))

    return staged


def _move(path: str, folder: str) -> None:
    os.makedirs(folder, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = os.path.basename(path)
    shutil.move(path, os.path.join(folder, f"{stamp}-{base}"))


def process_file(conn, path: str, mode: str) -> tuple[int, int]:
    """Ucitava jedan fajl. Vraca (primenjeno, odbijeno)."""
    rows = read_rows(path)
    staged = rows_to_staging(rows)

    cursor = conn.cursor()
    cursor.execute(
        sql(conn,
            "DECLARE @id BIGINT; "
            "EXEC etl.usp_StartLoadRun @JobName='STOCK_IMPORT', "
            "@SourceName=%s, @LoadRunId=@id OUTPUT; SELECT @id"),
        (os.path.basename(path),),
    )
    load_run_id = int(cursor.fetchone()[0])

    cursor.execute("TRUNCATE TABLE staging.StockImport")

    if staged:
        placeholders = ", ".join(["%s"] * (len(STAGING_COLUMNS) + 1))
        insert_sql = sql(conn, (
            f"INSERT INTO staging.StockImport (LoadRunId, {', '.join(STAGING_COLUMNS)}) "
            f"VALUES ({placeholders})"
        ))
        cursor.executemany(insert_sql, [(load_run_id, *row) for row in staged])

    conn.commit()

    cursor.execute(
        sql(conn, "EXEC inv.usp_ImportStockStaging @LoadRunId = %s, @Mode = %s"),
        (load_run_id, mode),
    )
    summary = cursor.fetchone()
    conn.commit()

    read, applied, rejected = int(summary[0]), int(summary[1]), int(summary[2])
    logger.info("%s: pročitano=%s primenjeno=%s odbijeno=%s",
                os.path.basename(path), read, applied, rejected)
    return applied, rejected


def run(drop_folder: str, settings: ConnectionSettings, mode: str) -> int:
    if not os.path.isdir(drop_folder):
        logger.error("Drop folder ne postoji: %s", drop_folder)
        return 1

    candidates = sorted(
        os.path.join(drop_folder, name)
        for name in os.listdir(drop_folder)
        if os.path.splitext(name)[1].lower() in (".xlsx", ".xlsm", ".csv", ".tsv")
        and os.path.isfile(os.path.join(drop_folder, name))
    )

    if not candidates:
        logger.info("Nema novih fajlova u %s.", drop_folder)
        return 0

    total_applied = total_rejected = 0

    with connection(settings) as conn:
        for path in candidates:
            try:
                applied, rejected = process_file(conn, path, mode)
                total_applied += applied
                total_rejected += rejected
                _move(path, os.path.join(drop_folder, "archive"))
            except Exception as exc:
                # Jedan los fajl ne sme da zaustavi ostale.
                logger.error("Neuspeh na fajlu %s: %s", path, exc)
                _move(path, os.path.join(drop_folder, "failed"))

    logger.info("Ukupno: primenjeno=%s odbijeno=%s", total_applied, total_rejected)
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Uvezi zalihe iz drop foldera u SQL Server.")
    parser.add_argument("--database", default=None)
    parser.add_argument("--drop-folder", default=os.environ.get("CARDTRACKER_DROP_FOLDER", "drop"))
    parser.add_argument("--mode", choices=["DELTA", "ABSOLUTE"], default="DELTA",
                        help="DELTA dodaje na postojeće stanje, ABSOLUTE ga postavlja (popis)")
    parser.add_argument("--verbose", action="store_true")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(name)s: %(message)s",
    )

    settings = ConnectionSettings.from_env(database=args.database)
    return run(args.drop_folder, settings, args.mode)


if __name__ == "__main__":
    sys.exit(main())
